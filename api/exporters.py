# -*- coding: utf-8 -*-
"""Excel 导出构造：单份（发票横表 / 报价单+托书竖排+费用明细块）与批量（横排大表+明细第二表）。

端点（server.py）只负责 HTTP 与文件名，构造逻辑全部在这里。
"""
from typing import Any, Dict, List

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment

from schemas import DOC_TYPES

# 发票 15 字段的中文表头（导出横表行 1 用；列顺序以 schema 定义为准 = TMS 导入恒定）
INVOICE_FIELD_LABELS = {
    "invoice_type": "发票类型", "invoice_number": "发票号", "invoice_date": "开票日期",
    "seller_name": "开票方名称", "seller_tax_id": "开票方税号", "buyer_name": "受票方名称",
    "buyer_tax_id": "受票方税号", "goods_description": "货物描述", "hs_code": "HS编码",
    "amount_subtotal": "金额小计", "tax_rate": "税率", "tax_amount": "税额",
    "amount_total": "总金额", "currency": "币种", "incoterm": "贸易术语",
}
STATUS_ZH = {"high": "高", "medium": "中", "low": "低"}

# 三类单据字段中文表头（批量导出表头行 1 用，与前端 fieldLabel 同词）。
# 旧映射键名（invoice_no/total_charge/pol/pod…）与现行 schema 永不命中，批量表头一直裸英文 key。
FIELD_LABELS_ZH = {
    **INVOICE_FIELD_LABELS,
    "quote_type": "报价单类型", "quote_number": "报价单号", "quote_date": "报价日期",
    "valid_until": "有效期至", "forwarder_name": "货代公司", "client_name": "客户公司",
    "port_of_loading": "起运港", "port_of_discharge": "目的港", "transport_mode": "运输方式",
    "cargo_description": "货物描述", "packages": "件数", "gross_weight": "毛重(kg)",
    "volume": "体积(m³)", "chargeable_weight": "计费重(kg)", "transit_days": "运输时效",
    "booking_number": "托书号", "booking_date": "托书日期", "shipper_name": "发货人",
    "shipper_address": "发货人地址", "vessel_name": "船名", "voyage_number": "航次",
    "place_of_delivery": "交货地", "etd": "预计开船(ETD)", "eta": "预计到港(ETA)",
    "container_type": "箱型", "container_count": "箱量", "payment_terms": "付款条款",
    "remarks": "备注",
}


def build_single_workbook(doc_type: str, extract_result: Dict[str, Any], filename: str = "") -> Workbook:
    """单份导出：
    发票 → 转置横表（一行一票，TMS 直接导入格式；列顺序按 schema 固定 15 列）；
    报价单/托书 → 竖排字段表（字段/值/置信度/状态）；
    报价单另附费用明细块（charge_items 逐行 + 合计 + 置信度对照）。"""
    wb = Workbook()
    ws = wb.active
    ws.title = DOC_TYPES.get(doc_type, {}).get("sheet_title", "Extract Result")

    if doc_type == "invoice":
        # ===== 发票横表：一行一票 =====
        props = DOC_TYPES["invoice"]["schema"]["properties"]
        keys = list(props.keys())  # schema 定义顺序 = 列顺序，恒 15 列

        def _cells(k):
            f = extract_result.get(k)
            if not isinstance(f, dict) or "value" not in f:
                return ("" if f is None else f, "", "")
            conf = f.get("confidence")
            return (f.get("value", ""),
                    round(conf * 100, 1) if isinstance(conf, (int, float)) else "",
                    STATUS_ZH.get(f.get("status"), ""))

        triples = [_cells(k) for k in keys]
        ws.append(["文件名"] + [INVOICE_FIELD_LABELS.get(k, k) for k in keys])   # 行1：中文表头
        ws.append(["filename"] + keys)                                           # 行2：字段 key（TMS 映射用）
        ws.append([filename or "invoice"] + [t[0] for t in triples])             # 行3：值（一行一票）
        ws.append(["置信度%"] + [t[1] for t in triples])                          # 行4：置信度
        ws.append(["状态"] + [t[2] for t in triples])                             # 行5：状态（高/中/低）

        for row_idx in (1, 2):
            for cell in ws[row_idx]:
                cell.font = Font(bold=True, color="FFFFFF", size=10)
                cell.fill = PatternFill("solid", fgColor="1B2C3F")
                cell.alignment = Alignment(horizontal="center", vertical="center")
        ws.column_dimensions["A"].width = 30
        for i in range(2, len(keys) + 2):
            ws.column_dimensions[chr(64 + i) if i <= 26 else f"A{chr(64 + i - 26)}"].width = 16
        ws.freeze_panes = "C3"
        return wb

    # ===== 报价单/托书：竖排字段表 =====
    ws.title = "Extract Result"

    # 把带 confidence 的字段展平
    def _flatten(d: Dict[str, Any], prefix: str = "") -> List[tuple]:
        rows = []
        for k, v in d.items():
            key = f"{prefix}{k}" if not prefix else f"{prefix}.{k}"
            if isinstance(v, dict) and "value" in v:
                rows.append((key, v["value"], round(v.get("confidence", 0) * 100, 1), v.get("status", "unknown")))
            elif isinstance(v, dict):
                rows.extend(_flatten(v, key))
            elif isinstance(v, list):
                # 行项目：值区只放行数指引，逐行明细写在下方"费用明细"区块
                rows.append((key, f"→ 见下方费用明细（{len(v)} 行）", "", ""))
            else:
                rows.append((key, v, "", ""))
        return rows

    flat_rows = _flatten(extract_result)

    headers = ["字段", "值", "置信度%", "状态"]
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="2563EB")
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for row in flat_rows:
        ws.append(list(row))

    ws.column_dimensions["A"].width = 28
    ws.column_dimensions["B"].width = 42
    ws.column_dimensions["C"].width = 14
    ws.column_dimensions["D"].width = 12

    # ===== 报价单费用明细块：charge_items 逐行落表（TMS 要的费用数据本体）+ 合计 + 置信度对照 =====
    charge_items = extract_result.get("charge_items")
    if isinstance(charge_items, list) and charge_items:
        ITEM_COLS = [("description", "费用项"), ("unit", "单位"), ("rate", "单价"),
                     ("quantity", "数量"), ("amount", "金额"), ("currency", "币种")]

        def _cell_v(item, col):
            f = item.get(col) if isinstance(item, dict) else None
            return f.get("value", "") if isinstance(f, dict) else ("" if f is None else f)

        def _cell_c(item, col):
            f = item.get(col) if isinstance(item, dict) else None
            c = f.get("confidence") if isinstance(f, dict) else None
            return round(c * 100, 1) if isinstance(c, (int, float)) else ""

        start = ws.max_row + 2
        ws.cell(row=start, column=1, value="费用明细（charge_items）").font = Font(bold=True, size=11)
        hdr = start + 1
        for i, (_, zh) in enumerate(ITEM_COLS, start=1):
            c = ws.cell(row=hdr, column=i, value=zh)
            c.font = Font(bold=True, color="FFFFFF", size=10)
            c.fill = PatternFill("solid", fgColor="2563EB")
            c.alignment = Alignment(horizontal="center", vertical="center")
        r = hdr
        total = 0.0
        for item in charge_items:
            r += 1
            for i, (col, _) in enumerate(ITEM_COLS, start=1):
                ws.cell(row=r, column=i, value=_cell_v(item, col))
            try:
                total += float(str(_cell_v(item, "amount")).replace(",", "") or 0)
            except (TypeError, ValueError):
                pass
        r += 1
        ws.cell(row=r, column=4, value="合计 Total").font = Font(bold=True, size=10)
        ws.cell(row=r, column=5, value=round(total, 2)).font = Font(bold=True, size=10)

        # 置信度对照块：与上方明细逐行对应（明细字段同样参与置信度分流）
        r += 2
        ws.cell(row=r, column=1, value="费用明细 · 置信度%（与上方明细逐行对应）").font = Font(bold=True, size=10)
        hdr2 = r + 1
        for i, (_, zh) in enumerate(ITEM_COLS, start=1):
            c = ws.cell(row=hdr2, column=i, value=zh)
            c.font = Font(bold=True, color="FFFFFF", size=10)
            c.fill = PatternFill("solid", fgColor="1B2C3F")
            c.alignment = Alignment(horizontal="center", vertical="center")
        rr = hdr2
        for item in charge_items:
            rr += 1
            for i, (col, _) in enumerate(ITEM_COLS, start=1):
                ws.cell(row=rr, column=i, value=_cell_c(item, col))

    return wb


def build_batch_workbook(doc_type: str, items: List[Any]) -> Workbook:
    """批量导出：所有单据 → 一张横排大表（TMS 直接导入格式）；
    报价单另附费用明细第二张表（一行一费用项，文件名关联主表）。items 为 BatchExportItem 列表。"""
    # 步骤 1：收集所有字段名（所有单据的字段并集做表头）
    def _collect_fields(d: Dict[str, Any], prefix: str = "") -> List[str]:
        fields = []
        for k, v in d.items():
            key = f"{prefix}{k}" if not prefix else f"{prefix}.{k}"
            if isinstance(v, dict) and "value" in v:
                fields.append(key)
            elif isinstance(v, dict):
                fields.extend(_collect_fields(v, key))
        return fields

    all_fields = []
    seen = set()
    for item in items:
        for f in _collect_fields(item.extract_result):
            if f not in seen:
                seen.add(f)
                all_fields.append(f)

    def _flatten_values(d: Dict[str, Any]) -> Dict[str, Any]:
        """把 extract_result 展平成 {key1: val1, key2: val2}"""
        out = {}
        def _walk(node, prefix=""):
            for k, v in node.items():
                key = f"{prefix}{k}" if not prefix else f"{prefix}.{k}"
                if isinstance(v, dict) and "value" in v:
                    out[key] = v.get("value", "")
                elif isinstance(v, dict):
                    _walk(v, key)
        _walk(d)
        return out

    # 步骤 2：生成 Excel
    wb = Workbook()
    ws = wb.active
    sheet_title = DOC_TYPES.get(doc_type, {}).get("sheet_title", doc_type)
    ws.title = sheet_title

    # 表头行 1：中文 label（三类单据字段全集映射；命不中再退字段名末段）
    header_labels = ["文件名"] + [FIELD_LABELS_ZH.get(f, f.split(".")[-1]) for f in all_fields]
    ws.append(header_labels)

    # 表头行 2：字段 key（TMS 导入需要）
    header_keys = ["filename"] + all_fields
    ws.append(header_keys)

    # 设置表头样式
    for row_idx in (1, 2):
        for cell in ws[row_idx]:
            cell.font = Font(bold=True, color="FFFFFF", size=10)
            cell.fill = PatternFill("solid", fgColor="1B2C3F")
            cell.alignment = Alignment(horizontal="center", vertical="center")

    # 数据行
    for item in items:
        flat = _flatten_values(item.extract_result)
        row = [item.filename]
        for f in all_fields:
            val = flat.get(f, "")
            row.append(val)
        ws.append(row)

    # 列宽
    ws.column_dimensions["A"].width = 32
    for i in range(2, len(all_fields) + 2):
        ws.column_dimensions[chr(64 + i) if i <= 26 else f"A{chr(64+i-26)}"].width = 18

    # 冻结前两行
    ws.freeze_panes = "A3"

    # ===== 报价单批量：费用明细第二张表（一行一费用项，文件名关联主表——TMS 主表+明细两表结构） =====
    detail_rows = []
    for item in items:
        er = item.extract_result or {}
        if isinstance(er.get("charge_items"), list):
            for ci in er["charge_items"]:
                def _v(col):
                    f = ci.get(col) if isinstance(ci, dict) else None
                    return f.get("value", "") if isinstance(f, dict) else ("" if f is None else f)
                detail_rows.append(
                    [item.filename] + [_v(c) for c in ("description", "unit", "rate", "quantity", "amount", "currency")])
    if detail_rows:
        ws2 = wb.create_sheet("费用明细")
        ws2.append(["文件名", "费用项", "单位", "单价", "数量", "金额", "币种"])
        for cell in ws2[1]:
            cell.font = Font(bold=True, color="FFFFFF", size=10)
            cell.fill = PatternFill("solid", fgColor="1B2C3F")
            cell.alignment = Alignment(horizontal="center", vertical="center")
        for row in detail_rows:
            ws2.append(row)
        ws2.column_dimensions["A"].width = 32
        for _c in "BCDEFG":
            ws2.column_dimensions[_c].width = 16
        ws2.freeze_panes = "A2"

    return wb
